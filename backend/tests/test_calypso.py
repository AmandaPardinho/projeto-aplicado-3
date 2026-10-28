from datetime import datetime
from decimal import Decimal

import openpyxl
import pytest

from report_calypso import CASO_703_1208, gerar_report, para_strict
from schulz.calypso import RelatorioInvalidoError, _Leitor, ler_relatorio
from schulz.xlsx import abrir_workbook, eh_strict


def regras(exc_info):
    return [v.regra for v in exc_info.value.violacoes]


def test_le_o_caso_703_1208():
    relatorio = ler_relatorio(gerar_report())
    assert len(relatorio.caracteristicas) == 14
    assert relatorio.plano == "703.1208_TESTE"
    assert relatorio.tag_do_plano == "703.1208"
    assert relatorio.data_hora == datetime(2026, 8, 13, 16, 16, 26)
    assert relatorio.metrologista == "Tecnico Exemplo"
    assert relatorio.codigo_mmc == "500008"


def test_separa_as_duas_passagens():
    c = ler_relatorio(gerar_report()).caracteristicas[4]  # 62P_DIAMETRO 2
    assert (c.medicao_p1, c.medicao_p2) == (Decimal("6.0781314"), Decimal("6.0780314"))
    assert (c.celula_p1, c.celula_p2) == ("B18", "B33")


def test_numeros_em_decimal_exato_rnf10():
    c = ler_relatorio(gerar_report()).caracteristicas[0]
    assert c.medicao_p1 == Decimal("0.002167")
    assert str(c.medicao_p1) == "0.002167"


def test_rh15_separa_cota_e_grandeza():
    caracteristicas = ler_relatorio(gerar_report()).caracteristicas
    assert (caracteristicas[2].cota_desenho, caracteristicas[2].grandeza) == ("62P", "DIÂMETRO 1")
    assert (caracteristicas[0].cota_desenho, caracteristicas[0].grandeza) == (None, "PLANICIDADE_REF")


def test_rf07_numero_variavel_de_caracteristicas():
    assert len(ler_relatorio(gerar_report(CASO_703_1208[:3])).caracteristicas) == 3
    muitas = [(f"{i}P_DIÂMETRO 1", 6.0, 6.0, 6.0, 0.01, -0.01) for i in range(35)]
    assert len(ler_relatorio(gerar_report(muitas)).caracteristicas) == 35


def test_rnf02_le_arquivo_strict():
    strict = para_strict(gerar_report())
    assert eh_strict(strict)
    assert abrir_workbook(strict).sheetnames == ["Report"]
    assert len(ler_relatorio(strict).caracteristicas) == 14


def test_rnf02_openpyxl_sozinho_nao_le_strict():
    import io
    import warnings
    strict = para_strict(gerar_report())
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        try:
            abas = openpyxl.load_workbook(io.BytesIO(strict)).sheetnames
        except Exception:
            abas = []
    assert abas == []


def test_rh12_ignora_abas_mortas():
    relatorio = ler_relatorio(gerar_report(abas_extras=("ViK", "KVi")))
    assert len(relatorio.caracteristicas) == 14


def test_rn01_recusa_uma_passagem_so():
    with pytest.raises(RelatorioInvalidoError) as exc:
        ler_relatorio(gerar_report(passagens=1))
    assert regras(exc) == ["RN01"]


def test_rn01_recusa_caracteristica_sem_segunda_passagem():
    with pytest.raises(RelatorioInvalidoError) as exc:
        ler_relatorio(gerar_report(substituir={(2, 3, 0): "OUTRA"}))
    assert "RN01" in regras(exc)


def test_rh16_erro_de_formula_vai_para_quarentena_com_a_celula():
    with pytest.raises(RelatorioInvalidoError) as exc:
        ler_relatorio(gerar_report(substituir={(1, 2, 1): "#VALUE!"}))
    violacao = next(v for v in exc.value.violacoes if v.regra == "RH16")
    assert violacao.celula == "B16"


def test_rn08_tolerancia_vazia():
    with pytest.raises(RelatorioInvalidoError) as exc:
        ler_relatorio(gerar_report(substituir={(1, 2, 3): "----", (2, 2, 3): "----"}))
    assert "RN08" in regras(exc)


def test_rh17_caracteristica_repetida_na_passagem():
    repetidas = CASO_703_1208[:2] + CASO_703_1208[:1]
    with pytest.raises(RelatorioInvalidoError) as exc:
        ler_relatorio(gerar_report(repetidas))
    assert "RH17" in regras(exc)


def test_rf02_recusa_planilha_sem_report():
    wb = openpyxl.Workbook()
    wb.active.title = "Registro"
    import io
    saida = io.BytesIO()
    wb.save(saida)
    with pytest.raises(RelatorioInvalidoError) as exc:
        ler_relatorio(saida.getvalue())
    assert regras(exc) == ["RF02"]


def test_rnf02_recusa_arquivo_que_nao_e_excel():
    with pytest.raises(RelatorioInvalidoError) as exc:
        ler_relatorio(b"isto nao e um xlsx")
    assert regras(exc) == ["RNF02"]


def test_rh13_remove_unidade_e_virgula():
    leitor = _Leitor()
    assert leitor.decimal("20,4°", "H5") == Decimal("20.4")
    assert leitor.transformacoes[0]["regra"] == "RH13"


def test_rh01_sentinelas_viram_nulo():
    leitor = _Leitor()
    assert leitor.texto("----", "B44") is None
    assert leitor.texto("???", "B45") is None
    assert leitor.texto("  Usinagem I ", "B24") == "Usinagem I"


def test_rh03_e_rh14_datas():
    leitor = _Leitor()
    assert leitor.data(46247, "D5") == datetime(2026, 8, 13)
    assert leitor.data("2026-08-26T16:47:10.00000000465661175", "C5") == \
        datetime(2026, 8, 26, 16, 47, 10)
    assert leitor.hora("04:17:00.000000000000375", "C4").isoformat() == "04:17:00"
