"""
Gera relatorios no layout do "Calypso Measuring Result" (aba Report).

Os valores do caso 703.1208 sao os do Registro_MMC.xlsx real, inclusive a
diferenca entre passagens do 106P_POSICAO 1 (0,1326897 x 0,1316897), que e a
maior dispersao e entra no orcamento GUM. O nome do metrologista e ficticio.
"""
from __future__ import annotations

import io
from datetime import datetime, time

import openpyxl

# (caracteristica, passagem 1, passagem 2, nominal, tol. superior, tol. inferior)
CASO_703_1208 = [
    ("PLANICIDADE_REF", 0.002167, 0.002267, 0, 0.05, 0),
    ("ENTRE CENTROS_REF", 0.0274147, 0.0274147, 0, 0.04, 0),
    ("62P_DIÂMETRO 1", 6.0787295, 6.0787295, 6.075, 0.01, -0.01),
    ("62P_POSIÇÃO 1", 0.4295522, 0.4295522, 0, 0.1, 0),
    ("62P_DIÂMETRO 2", 6.0781314, 6.0780314, 6.075, 0.01, -0.01),
    ("62P_POSIÇÃO 2", 0.2446114, 0.2446114, 0, 0.1, 0),
    ("13P_DIÂMETRO 1", 16.9684482, 16.9684482, 16.964, 0.01, -0.01),
    ("13P_POSIÇÃO 1", 0.9590437, 0.9590437, 0, 0.1, 0),
    ("13P_DIÂMETRO 2", 16.9698259, 16.9698259, 16.964, 0.01, -0.01),
    ("13P_POSIÇÃO 2", 0.2355755, 0.2355755, 0, 0.1, 0),
    ("106P_DIÂMETRO 1", 6.3463135, 6.3463135, 6.345, 0.01, -0.01),
    ("106P_POSIÇÃO 1", 0.1326897, 0.1316897, 0, 0.1, 0),
    ("106P_DIÂMETRO 2", 6.3485617, 6.3485617, 6.345, 0.01, -0.01),
    ("106P_POSIÇÃO 2", 0.867758, 0.867758, 0, 0.1, 0),
]


def gerar_report(caracteristicas=CASO_703_1208, plano="703.1208_TESTE",
                 metrologista="Tecnico Exemplo", codigo_mmc=500008,
                 data=datetime(2026, 8, 13), hora=time(16, 16, 26),
                 passagens=2, abas_extras=(), substituir=None) -> bytes:
    """substituir: {(passagem, indice, coluna): valor} para injetar defeitos.

    coluna: 0 nome, 1 medido, 2 nominal, 3 tol. sup., 4 tol. inf.
    """
    substituir = substituir or {}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["B2"] = "Calypso Measuring Result"
    for coluna, rotulo, valor in (("B", "Measurement Plan", plano), ("D", "Date", data),
                                  ("F", "Order", None)):
        ws[f"{coluna}4"], ws[f"{coluna}5"] = rotulo, valor
    for coluna, rotulo, valor in (("B", "Drawing No.", None), ("D", "Time", hora),
                                  ("F", "Part No.", 27)):
        ws[f"{coluna}7"], ws[f"{coluna}8"] = rotulo, valor
    for coluna, rotulo, valor in (("B", "Metrologista", metrologista), ("D", "MMC", codigo_mmc),
                                  ("F", "Fornecedor", "CDM")):
        ws[f"{coluna}10"], ws[f"{coluna}11"] = rotulo, valor
    for i, rotulo in enumerate(("Characteristic", "Actual", "Nominal", "Upper Tol",
                                "Lower Tol", "Deviation"), start=1):
        ws.cell(row=13, column=i, value=rotulo)

    linha = 14
    for passagem in range(1, passagens + 1):
        for indice, (nome, p1, p2, nominal, tol_sup, tol_inf) in enumerate(caracteristicas):
            medido = p1 if passagem == 1 else p2
            valores = [nome, medido, nominal, tol_sup, tol_inf]
            for coluna in range(5):
                valores[coluna] = substituir.get((passagem, indice, coluna), valores[coluna])
            for coluna, valor in enumerate(valores, start=1):
                ws.cell(row=linha, column=coluna, value=valor)
            if isinstance(valores[1], (int, float)) and isinstance(valores[2], (int, float)):
                ws.cell(row=linha, column=6, value=round(valores[1] - valores[2], 7))
            linha += 1
        linha += 1  # linha vazia entre as passagens

    for nome in abas_extras:
        wb.create_sheet(nome)["A1"] = "#NAME?"
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def para_strict(conteudo: bytes) -> bytes:
    """Converte um xlsx Transitional em Strict (o inverso do reparo)."""
    import zipfile

    from schulz.xlsx import NAMESPACES_STRICT

    origem = zipfile.ZipFile(io.BytesIO(conteudo))
    saida = io.BytesIO()
    with zipfile.ZipFile(saida, "w", zipfile.ZIP_DEFLATED) as destino:
        for item in origem.infolist():
            dados = origem.read(item.filename)
            if item.filename.endswith((".xml", ".rels")):
                for strict, transitional in NAMESPACES_STRICT.items():
                    dados = dados.replace(transitional, strict)
            destino.writestr(item, dados)
    return saida.getvalue()
