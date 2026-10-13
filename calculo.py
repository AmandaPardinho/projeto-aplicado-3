"""
Modulo de regras de negocio e calculos metrologicos do MVP (AV3).

Implementa o que o roteiro da AV3 pede, em cima da estrutura de dados ja
levantada no Tema 01 (aba "Importacao" / relatorio do Calypso):

    - Media das leituras de cada caracteristica (RN01: sempre duas medicoes)
    - Erro sistematico (media - valor nominal)
    - Validacao por ponto: erro vs. Limite de Erro Admissivel (LEA),
      aqui tratado como a tolerancia (+TOL / -TOL) de cada caracteristica
      (erro sistematico assinado dentro de -TOL..+TOL; equivale ao erro
      absoluto <= LEA quando as tolerancias sao simetricas)
    - Parecer final: APROVADO se todas as caracteristicas estiverem dentro
      da tolerancia, REPROVADO se qualquer uma estiver fora

Mantém as regras de higienizacao RH01/RH04/RH09 do Tema 01 na leitura:
nulos textuais, espacos em branco e linhas de caracteristica vazias sao
descartados antes de qualquer calculo. Ja uma caracteristica nomeada com
campo numerico ausente ou invalido NAO e descartada em silencio: a planilha
e recusada, indicando as linhas a corrigir, para nunca emitir parecer sobre
dados parciais.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Optional

import openpyxl

# Nomes de cabecalho aceitos (normalizados em minusculo, sem acento) para
# cada coluna. O relatorio real do Calypso usa "DIMENSION"/"Parametro
# avaliado"; deixamos sinonimos comuns tambem, para o parser nao travar se
# outra equipe/planilha usar um rotulo ligeiramente diferente.
HEADER_SYNONYMS = {
    "caracteristica": {"dimension", "parametro avaliado", "caracteristica", "característica"},
    "nominal": {"nominal"},
    "tol_mais": {"+tol", "tol +", "tol+"},
    "tol_menos": {"-tol", "tol -", "tol-"},
    "medida1": {"meas 1", "meas1", "medida 1", "leitura 1"},
    "medida2": {"meas 2", "meas2", "medida 2", "leitura 2"},
}

# Marcadores que RH01 trata como nulo textual
NULOS_TEXTUAIS = {"----", "-----", "???", ""}


class PlanilhaInvalidaError(ValueError):
    """Levantado quando a planilha nao tem a estrutura minima esperada."""


@dataclass
class ResultadoCaracteristica:
    nome: str
    nominal: Decimal
    tol_mais: Decimal
    tol_menos: Decimal
    medida1: Decimal
    medida2: Decimal
    media: Decimal = field(init=False)
    erro_sistematico: Decimal = field(init=False)
    status: str = field(init=False)

    def __post_init__(self) -> None:
        self.media = (self.medida1 + self.medida2) / 2
        self.erro_sistematico = self.media - self.nominal
        dentro = self.tol_menos <= self.erro_sistematico <= self.tol_mais
        self.status = "DENTRO_DA_TOLERANCIA" if dentro else "FORA_DA_TOLERANCIA"


@dataclass
class ResultadoCalibracao:
    caracteristicas: list[ResultadoCaracteristica]
    parecer: str = field(init=False)

    def __post_init__(self) -> None:
        reprovadas = [c for c in self.caracteristicas if c.status == "FORA_DA_TOLERANCIA"]
        self.parecer = "REPROVADO" if reprovadas else "APROVADO"

    @property
    def total_reprovadas(self) -> int:
        return sum(1 for c in self.caracteristicas if c.status == "FORA_DA_TOLERANCIA")


def _limpar_texto(valor) -> Optional[str]:
    """RH01/RH04: nulos textuais e espacos nas pontas viram None/trim."""
    if valor is None:
        return None
    texto = str(valor).strip()
    return None if texto in NULOS_TEXTUAIS else texto


def _normalizar_cabecalho(valor) -> Optional[str]:
    texto = _limpar_texto(valor)
    return texto.lower() if texto else None


def _para_decimal(valor) -> Optional[Decimal]:
    """RH13: aceita numero, ou texto com virgula/ponto decimal e remove
    simbolos de unidade residuais antes de converter."""
    if valor is None:
        return None
    if isinstance(valor, (int, float, Decimal)):
        return Decimal(str(valor))
    texto = _limpar_texto(valor)
    if texto is None:
        return None
    texto = texto.replace(",", ".")
    texto = "".join(ch for ch in texto if ch.isdigit() or ch in ".-")
    if texto in ("", "-", "."):
        return None
    try:
        return Decimal(texto)
    except InvalidOperation:
        return None


def _localizar_cabecalho(planilha):
    """Varre as primeiras linhas procurando a linha de cabecalho da tabela
    de caracteristicas (RH09: estruturar antes de limpar)."""
    for linha in planilha.iter_rows(min_row=1, max_row=min(30, planilha.max_row)):
        indices = {}
        for celula in linha:
            cab = _normalizar_cabecalho(celula.value)
            if not cab:
                continue
            for campo, sinonimos in HEADER_SYNONYMS.items():
                if cab in sinonimos:
                    indices[campo] = celula.column
        if {"caracteristica", "nominal", "tol_mais", "tol_menos", "medida1", "medida2"} <= indices.keys():
            return linha[0].row, indices
    raise PlanilhaInvalidaError(
        "Nao foi possivel localizar o cabecalho da tabela de caracteristicas "
        "(esperado: DIMENSION/Parametro avaliado, NOMINAL, +TOL, -TOL, MEAS 1, MEAS 2)."
    )


def processar_planilha(caminho_arquivo: str) -> ResultadoCalibracao:
    """Le o arquivo .xlsx enviado e devolve o resultado da calibracao
    (lista de caracteristicas calculadas + parecer final)."""
    try:
        workbook = openpyxl.load_workbook(caminho_arquivo, data_only=True)
    except Exception as exc:  # arquivo corrompido, formato nao suportado, etc.
        raise PlanilhaInvalidaError(f"Nao foi possivel abrir o arquivo: {exc}") from exc

    planilha = workbook.active
    linha_cabecalho, indices = _localizar_cabecalho(planilha)

    resultados: list[ResultadoCaracteristica] = []
    problemas: list[str] = []
    for linha in planilha.iter_rows(min_row=linha_cabecalho + 1, max_row=planilha.max_row):
        nome = _limpar_texto(linha[indices["caracteristica"] - 1].value)
        # RH09: linha de caracteristica vazia e descartada antes de qualquer calculo
        if not nome:
            continue

        nominal = _para_decimal(linha[indices["nominal"] - 1].value)
        tol_mais = _para_decimal(linha[indices["tol_mais"] - 1].value)
        tol_menos = _para_decimal(linha[indices["tol_menos"] - 1].value)
        medida1 = _para_decimal(linha[indices["medida1"] - 1].value)
        medida2 = _para_decimal(linha[indices["medida2"] - 1].value)

        valores = {
            "NOMINAL": nominal,
            "+TOL": tol_mais,
            "-TOL": tol_menos,
            "MEAS 1": medida1,
            "MEAS 2": medida2,
        }
        ausentes = [rotulo for rotulo, valor in valores.items() if valor is None]
        if ausentes:
            # Nao calcula com dados parciais: acumula para recusar a planilha
            # inteira, evitando um parecer que ignore uma caracteristica.
            problemas.append(f"linha {linha[0].row} ('{nome}'): {', '.join(ausentes)}")
            continue

        resultados.append(
            ResultadoCaracteristica(
                nome=nome,
                nominal=nominal,
                tol_mais=tol_mais,
                tol_menos=tol_menos,
                medida1=medida1,
                medida2=medida2,
            )
        )

    if problemas:
        raise PlanilhaInvalidaError(
            "Planilha incompleta: corrija os campos ausentes ou invalidos e envie "
            "novamente. " + "; ".join(problemas) + "."
        )

    if not resultados:
        raise PlanilhaInvalidaError(
            "Nenhuma linha de caracteristica valida foi encontrada na planilha."
        )

    return ResultadoCalibracao(caracteristicas=resultados)
